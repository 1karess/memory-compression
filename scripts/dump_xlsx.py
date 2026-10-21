"""Dump every sheet of an xlsx (values, formulas, merges, fills, fonts, widths) for inspection."""
import sys
from openpyxl import load_workbook

path = sys.argv[1]
wb = load_workbook(path)
wbv = load_workbook(path, data_only=True)
for ws in wb.worksheets:
    wv = wbv[ws.title]
    print(f"=== SHEET {ws.title!r} dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column} freeze={ws.freeze_panes}")
    print("merged:", [str(m) for m in ws.merged_cells.ranges][:80])
    print("col widths:", {k: round(v.width, 1) for k, v in ws.column_dimensions.items() if v.width})
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            v = wv[c.coordinate].value
            fill = c.fill.fgColor.rgb if c.fill and c.fill.fill_type else None
            font = f"{c.font.name},{c.font.sz},{'B' if c.font.b else ''},{c.font.color.rgb if c.font.color else None}"
            dv = f" cached={v!r}" if isinstance(c.value, str) and c.value.startswith("=") else ""
            print(f"{c.coordinate}: {c.value!r}{dv} | fill={fill} font={font} fmt={c.number_format}")
    if ws.data_validations.dataValidation:
        print("validations:", [(d.sqref, d.formula1) for d in ws.data_validations.dataValidation])
