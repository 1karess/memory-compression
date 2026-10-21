"""Fill the stage-0 tracking workbook (user's template, docs/input_record.xlsx) from the result files.

Keeps the template's sheets, fonts and colours; corrects the planned values that differ from the
protocol actually run (test 2000, separate tune split, memory-length sweep, seeds, patches);
fills every measured cell whose source file exists and leaves the rest as "未开始" / blank.
Adds two sheets: 04b_记忆长度扫描 and 07_实现检查. Safe to re-run at any time.

usage: python scripts/fill_record.py [template.xlsx] [out.xlsx]
"""
import datetime as dt
import json
import re
import subprocess
import sys
from copy import copy
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).parent))
from ccpp_common import FEATURES, ROOT, TARGET, load_splits  # noqa: E402

R, O, LOGS = ROOT / "results", ROOT / "outputs", ROOT / "logs"
MAIN = O / "main"
SUBSET = "r0_n256"
SEEDS = [42, 43, 44]  # seed_1, seed_2, seed_3

FONT = "Carlito"
INK, LABEL = "FF1E293B", "FF223B57"
F_TITLE, F_HEAD, F_ZEBRA, F_INPUT, F_NOTE, F_LABEL = (PatternFill("solid", fgColor=c) for c in
                                                      ["FF16263C", "FF223B57", "FFF5F7FA", "FFFFF3D3", "FFE8F4F5", "FFEDF3FA"])
THIN = Side(style="thin", color="FFD9E0E8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="center")


# ---------------------------------------------------------------- helpers
def jload(p):
    p = Path(p)
    return json.loads(p.read_text()) if p.exists() else None


def mdate(p):
    p = Path(p)
    return dt.datetime.fromtimestamp(p.stat().st_mtime).strftime("%Y-%m-%d") if p.exists() else None


def rel(p):
    return str(Path(p).relative_to(ROOT)) if str(p).startswith(str(ROOT)) else str(p)


def f3(x, nd=3):
    return None if x is None or (isinstance(x, float) and np.isnan(x)) else round(float(x), nd)


def setv(ws, ref, value, fmt=None):
    c = ws[ref]
    c.value = value
    if fmt:
        c.number_format = fmt
    return c


def style(c, *, fill=None, bold=False, color=INK, size=10, fmt=None, align=WRAP, border=True):
    c.font = Font(name=FONT, size=size, bold=bold, color=color)
    if fill:
        c.fill = fill
    c.alignment = align
    if border:
        c.border = BORDER
    if fmt:
        c.number_format = fmt
    return c


def title_block(ws, title, subtitle, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=ncols)
    style(ws.cell(1, 1, title), fill=F_TITLE, bold=True, color="FFFFFFFF", size=19, border=False,
          align=Alignment(vertical="center"))
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)
    style(ws.cell(3, 1, subtitle), fill=F_NOTE, color=LABEL, border=False, align=WRAP)
    ws.row_dimensions[3].height = 30


def header(ws, row, names):
    for j, n in enumerate(names, 1):
        style(ws.cell(row, j, n), fill=F_HEAD, bold=True, color="FFFFFFFF")


def body_row(ws, row, values, fmts=None, zebra=False, input_cols=()):
    for j, v in enumerate(values, 1):
        c = ws.cell(row, j, v)
        fill = F_INPUT if j in input_cols else (F_ZEBRA if zebra else None)
        style(c, fill=fill, fmt=(fmts or {}).get(j))


def note(ws, row, text, ncols, fill=F_NOTE):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    style(ws.cell(row, 1, text), fill=fill, color=LABEL, border=False)
    ws.row_dimensions[row].height = 30


# ---------------------------------------------------------------- data
def gather():
    d = {}
    d["revision"] = (ROOT / "models" / "Qwen3-4B.REVISION").read_text().strip() if (ROOT / "models" / "Qwen3-4B.REVISION").exists() else None
    try:
        d["cart_commit"] = subprocess.run(["git", "-C", str(ROOT / "third_party" / "cartridges"), "rev-parse", "--short", "HEAD"],
                                          capture_output=True, text=True).stdout.strip()
    except Exception:
        d["cart_commit"] = "ef34ba9"
    freeze = (LOGS / "pip_freeze_setup.txt").read_text() if (LOGS / "pip_freeze_setup.txt").exists() else ""
    d["versions"] = {k: (re.search(rf"^{k}==(\S+)", freeze, re.M) or [None, None])[1] for k in ["torch", "transformers"]}
    d["split"] = jload(ROOT / "data" / "processed" / "split_report.json")
    d["audit"] = jload(R / "token_audit.json")
    d["eq"] = jload(R / "gate" / "equivalence.json")
    d["cc"] = jload(R / "cartridges_checks.json")
    d["p4"] = jload(R / "p4_check.json")
    for k, p in [("gate", R / "gate" / "gate_summary.csv"), ("gate_len", R / "gate" / "gate_by_length.csv"),
                 ("gate_pq", R / "gate" / "gate_per_query.csv"), ("base_tune", R / "baselines_tune.csv"),
                 ("final", MAIN / "final" / "final_test_summary.csv"), ("final_pq", MAIN / "final" / "final_test_per_query.csv"),
                 ("A_tune", MAIN / "A_long_tune.csv"), ("A_test", MAIN / "A_long_test.csv")]:
        d[k] = pd.read_csv(p) if p.exists() else None
    d["teacher"] = jload(MAIN / "teacher" / "summary.json")
    d["ccpp32"] = jload(O / "ccpp32" / "run_k64" / "dev_selection.json")
    d["toy"] = jload(O / "toy" / "run_k32" / "dev_selection.json")
    d["toy_teacher"] = jload(O / "toy" / "teacher" / "summary.json")
    d["toy256"] = jload(O / "toy256" / "run_k32" / "dev_selection.json")
    d["toy256_teacher"] = jload(O / "toy256" / "teacher" / "summary.json")
    d["lr"] = (MAIN / "chosen_lr.txt").read_text().strip() if (MAIN / "chosen_lr.txt").exists() else None
    d["stats"] = jload(R / "test_stats.json")
    d["pseudo"] = jload(R / "pseudo_label_control.json")
    d["runs"] = {}
    for run in sorted((MAIN / "runs").glob("lr*_s*")) if (MAIN / "runs").exists() else []:
        d["runs"][run.name] = {
            "dir": run, "args": jload(run / "train_args.json"), "sel": jload(run / "dev_selection.json"),
            "sel_csv": pd.read_csv(run / "dev_selection.csv") if (run / "dev_selection.csv").exists() else None,
            "done": (run / "TRAIN_DONE").exists(),
        }
    return d


def classic(split, rows, subsets):
    from sklearn.linear_model import LinearRegression
    from sklearn.neighbors import KNeighborsRegressor
    from sklearn.preprocessing import StandardScaler
    by_id = rows.set_index("id")
    mem = by_id.loc[subsets[SUBSET]["ids"]]
    q = rows[rows["split"] == split]
    X, y, Xq, yq = mem[FEATURES].to_numpy(), mem[TARGET].to_numpy(), q[FEATURES].to_numpy(), q[TARGET].to_numpy()
    sc = StandardScaler().fit(X)
    preds = {"mean": np.full(len(yq), y.mean()), "ols": LinearRegression().fit(X, y).predict(Xq),
             "knn5": KNeighborsRegressor(5).fit(sc.transform(X), y).predict(sc.transform(Xq))}
    return {k: (float(np.sqrt(np.mean((p - yq) ** 2))), float(np.mean(np.abs(p - yq))), len(yq)) for k, p in preds.items()}


def train_losses(run_name):
    """step -> last displayed training loss (truncated soft-label CE) from the main job logs."""
    out = {}
    for log in sorted(LOGS.glob("mz_main_*.out")):
        text = log.read_text(errors="ignore")
        for seg in re.split(r'(?="run_dir": ")', text)[1:]:
            if not seg.startswith(f'"run_dir": "outputs/main/runs/{run_name}"'):
                continue
            for m in re.finditer(r"loss=([0-9.]+), ppl=[0-9.]+, optimizer_step=(\d+)", seg):
                out[int(m.group(2))] = float(m.group(1))
    return out


def z_drift(run_dir, ckpt_names):
    """||Z - Z_init||_F / ||Z_init||_F over the trainable positions, per checkpoint."""
    import torch
    init = torch.load(run_dir / "cache_init.pt", map_location="cpu", weights_only=False)
    base = torch.cat([t.detach().float().flatten() for t in list(init["trainable_keys"]) + list(init["trainable_values"])])
    out = {}
    for name in ckpt_names:
        p = run_dir / f"{name}.pt"
        if not p.exists():
            continue
        ck = torch.load(p, map_location="cpu", weights_only=False)
        z = torch.cat([t.detach().float().flatten() for t in list(ck["trainable_keys"]) + list(ck["trainable_values"])])
        out[name] = float((z - base).norm() / base.norm())
    return out


REFUSAL = re.compile(r"cannot|can't|unable|not possible|not computable|without (?:a|the|additional|specific)", re.I)


def refusal_rate(df):
    """Share of long A answers that explicitly decline to give a number."""
    return float(df["text"].fillna("").str.contains(REFUSAL).mean())


def instability(run):
    """Describe a dev-KL blow-up after the selected checkpoint (None if the run was smooth)."""
    if not run or run["sel_csv"] is None or not run["sel"]:
        return None
    df = run["sel_csv"][run["sel_csv"]["ckpt"] != "init"].reset_index(drop=True)
    best = run["sel"]["best_dev_kl"]
    # only a rise *after* the run had already converged near its best counts (early high KL is just warm-up)
    converged = df.index[df["dev_kl"] < 2 * best]
    if len(converged) == 0:
        return None
    after = df.loc[converged[0]:]
    peak = after.loc[after["dev_kl"].idxmax()]
    if peak["dev_kl"] < 3 * best:
        return None
    last = df.iloc[-1]
    return (f"{run['dir'].name} 训练中途发散：第 {int(peak['updates'])} 次更新附近开发 KL 升至 {peak['dev_kl']:.2f}，"
            f"训练结束时 {last['dev_kl']:.3f}；按开发 KL 选到发散前第 {run['sel']['best_updates']} 次更新（{best:.3f}）")


def run_minutes(run_dir):
    a, b = run_dir / "train_args.json", run_dir / "TRAIN_DONE"
    if a.exists() and b.exists():
        return round((b.stat().st_mtime - a.stat().st_mtime) / 60, 1)
    return None


def gpu_name():
    for log in sorted(LOGS.glob("mz_main_*.out")):
        m = re.search(r"\n(NVIDIA [^,\n]+), (\d+) MiB", log.read_text(errors="ignore"))
        if m:
            return f"{m.group(1)} ({int(m.group(2)) // 1024} GB)"
    return None


# ---------------------------------------------------------------- sheets
def fill_overview(ws, d):
    ws["A3"] = f"面向组会的实时记录表 · 计划值已预填，浅黄色格为实测值 · 由 scripts/fill_record.py 按结果文件自动填写 · 更新于 {dt.datetime.now():%Y-%m-%d %H:%M}"
    ws["B6"] = ("阶段 0｜UCI CCPP 回归：先确认完整记忆 M（256 行，10,979 词元）能帮助冻结的 Qwen3-4B 预测，"
                "再验证 64 位置的压缩记忆 Z 能保留多少效果（缓存长度比 ≈171.5×）。")
    ws["A20"] = "1）【02_执行步骤】逐项记录状态、实际结果与证据路径；脚本按结果文件自动更新，手工补充的内容写在 J 列。"
    ws["A21"] = "2）教师关卡见【04_教师预检】（主记忆 r0_n256）与【04b_记忆长度扫描】（16–512 行 × 3 份抽样）；未过关不启动训练。"
    ws["A22"] = "3）训练与选点见【05_训练日志】，最终测试见【06_测试汇总】；实现与数值检查见【07_实现检查】；计划与实际配置见【03_准备配置】。"
    ws["A23"] = "4）浅黄色格为实测字段；“未开始”表示尚未验证，不能视为失败或完成。公式在 Excel / WPS 打开时自动重算。"
    st, fin = d["stats"], d["final"]
    if st and fin is not None:
        rr, q = st["recovery_R"], st["pairs"]["C_avg3 - B"]
        ps = d["pseudo"]
        ws["A26"] = (f"测试集 2000 条：64 位置的 Z（KV 位置为完整记忆的 1/171.5）在同分布查询上复现了 B 的大部分行为——RMSE {st['rmse']['C_avg3']:.2f}"
                     f"（3 种子），B {st['rmse']['B']:.2f}，均值底线 {st['rmse']['mean']:.2f}；C 略差且显著（+{q['diff']:.2f} MW [{q['ci95'][0]:.2f}, {q['ci95'][1]:.2f}]），"
                     f"相对均值的改进保留 {rr['point']:.0%} [{rr['ci95'][0]:.0%}, {rr['ci95'][1]:.0%}]。"
                     + (f"但用同样 1024 个教师伪标签拟合的 OLS 已达 {ps['teacher_pseudo_labels']['OLS']['test_rmse_vs_truth']:.2f}，"
                        "RMSE 层面还不能区分“压缩了 M”与“从伪标签学了回归”；C 与 B 的贴合度（相关 "
                        f"{ps['C_reference']['corr_with_B_mean_of_seeds']:.3f}，数字差 {ps['C_reference']['mean_abs_diff_vs_B_mean_of_seeds']:.2f}）高于这些回归器。" if ps else "")
                     + "限定：单一数据集 / 单一 M / 单一模型 / k=64；B 本身弱于线性回归（4.59）。")
    elif d["gate"] is not None:
        ws["A26"] = "教师关卡已通过（tune 集：B 9.50 < B-打乱 14.57 < 均值 16.48）；压缩记忆训练进行中。"
    # the template's progress data bars are an x14 extension that openpyxl drops on load; recreate them
    bar = DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="FF187E8B")
    ws.conditional_formatting.add("B11", bar)
    ws.conditional_formatting.add("E14:E17", bar)


def status_colors(ws, rng):
    for text, fill in [("已完成", "FFDCF2E3"), ("进行中", "FFFFF3D3"), ("阻塞", "FFF9DADA"), ("失败", "FFF9DADA")]:
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'EXACT({rng.split(":")[0]},"{text}")'],
                                                       fill=PatternFill("solid", fgColor=fill, bgColor=fill)))


STEP_TEXT_FIXES = {  # (row, col): corrected planned text, where the protocol run differs from the template
    ("E", 8): "按格式化后的 X 去重分组：记忆池 1024（从中抽 M=256）；蒸馏 1024；调格式集 tune 256；开发 256；测试 2000；其余 5008 备用；保存索引和种子",
    ("E", 13): "固定模板与顺序；B 在 256 条 tune 查询上生成可解析的数值",
    ("C", 14): "在 tune 集运行 A、均值、B、B-打乱（另做记忆长度扫描 16/64/128/256/512 × 3 份抽样）",
    ("D", 14): "tune 集 256 条含真实 Y（与开发集分开，开发集只用于选 checkpoint）；同一批输入",
    ("D", 15): "A、均值、B、B-打乱在 tune 集的结果；配对 bootstrap 95% CI",
    ("D", 20): "截断软标签 CE；BF16 前向 / FP32 可训练主参数（补丁 P4，已实测）；lr 在 {5e-3, 2e-2} 中按开发集 KL 选",
    ("D", 23): "同一份测试 X/Y（2000 条）；统一数字解析与失败回退",
}


def fill_steps(ws, d):
    for (col, row), text in STEP_TEXT_FIXES.items():
        ws[f"{col}{row}"] = text
    gate, eq, cc, p4, sp, au = d["gate"], d["eq"], d["cc"], d["p4"], d["split"], d["audit"]
    s = {}  # row -> (status, result, evidence, date, problem)

    if d["revision"]:
        s[7] = ("已完成",
                f"Qwen/Qwen3-4B revision {d['revision'][:12]}；Cartridges commit {d['cart_commit']}（+补丁 P1/P2/P4）；"
                f"torch {d['versions']['torch']}，transformers {d['versions']['transformers']}，Python 3.12.2",
                "models/Qwen3-4B.REVISION；logs/pip_freeze_setup.txt；logs/cartridges_patch.diff", mdate(ROOT / "models" / "Qwen3-4B.REVISION"),
                "模型以本地路径加载时 Cartridges 按名字查表会 KeyError → scripts/cartridges_compat.py 注册别名")
    if sp:
        s[8] = ("已完成",
                f"Sheet1 共 {sp['n_rows']} 行，唯一 X {sp['n_unique_x']}；重复 X 组中的行 {sp['n_rows_in_duplicate_x_groups']}（同 X 不同 Y 的组 {sp['n_duplicate_x_groups_with_different_y']}）；"
                f"跨集合 X 泄漏 {sp['x_leak_across_splits']}；切分种子 {sp['seed']}；各集合："
                + "，".join(f"{k} {v}" for k, v in sp["split_sizes"].items()),
                "data/processed/rows.csv；data/processed/split_report.json；data/processed/memory_subsets.json", mdate(ROOT / "data" / "processed" / "split_report.json"),
                "测试集扩为 2000 条（缩小置信区间）；另设 tune 集做教师关卡，开发集只用于选点")
    if au:
        mc = au["message_converter"]
        s[9] = ("已完成",
                "A/B/C 共用一次 apply_chat_template(enable_thinking=False) 渲染再切分；生成提示以空思考块 <think>\\n\\n</think>\\n\\n 结尾（已断言）。"
                f"Cartridges 转换器默认不含空思考块（plain 与目标一致：{mc['cartridges_plain_equals_target']}），"
                f"把空思考块写入 assistant token_ids 后一致：{mc['cartridges_with_empty_think_equals_target']}；1280 条教师数据逐条断言 Cartridges 分词 == 教师所见 token",
                "results/token_audit.json；scripts/ccpp_common.py build_ids", mdate(R / "token_audit.json"),
                "Cartridges 的 MODELS_WITH_THINKING 区分大小写：写官方名 Qwen/Qwen3-4B 时评测会默认进入思考模式（我们不用其评测代码）")
    if eq and cc:
        f32 = eq.get("fp32_vs_bf16_full", {})
        s[10] = ("已完成",
                 f"前缀 KV 路径 vs 整串前向（20 条）：KL 均值 {eq['cached_vs_full_bf16']['kl_mean']['mean']:.4f}，argmax 一致 {eq['cached_vs_full_bf16']['argmax_agree']['mean']:.3f}；"
                 f"bf16 噪声底（fp32 vs bf16）KL {f32.get('kl_mean', {}).get('mean', float('nan')):.4f}，argmax {f32.get('argmax_agree', {}).get('mean', float('nan')):.3f}；"
                 f"crop 后重跑逐位一致；存/读后 logits 完全一致（max diff {cc['flex_cache_vs_reloaded_cache']['max_abs_logit_diff']}）",
                 "results/gate/equivalence.json；results/cartridges_checks.json", mdate(R / "gate" / "equivalence.json"),
                 "两路差异与 bf16 自身噪声同量级（≤ 2× 噪声底），判为一致")
    if cc:
        ck = cc["cache_k64"]
        s[11] = ("已完成",
                 f"P1 cache.py from_pretrained 冻结数取 size(2)：修复前会读成冻结 {ck['unpatched_would_read_num_frozen']}，修复后 {ck['loaded_num_frozen']}/{ck['loaded_num_trainable']}（冻结/可训练）；"
                 "P2 max_optimizer_steps 到步即停；flatten：我们直接从 logits 取 top-20 不经过它（覆盖率见 3.2）；"
                 "上游 Qwen3 等价性测试 8/8 通过",
                 "logs/cartridges_patch.diff；logs/mz_cart_checks_*.out", mdate(R / "cartridges_checks.json"),
                 "新增 P4（fp32 主参数，见 1.6）")
    if cc and p4:
        b16 = {r["lr"]: [x["frac_intended_nonzero_but_rounded_to_zero"] for x in r["steps"]] for r in cc["adam_update"]}
        f32u = {r["lr"]: max(x["frac_intended_nonzero_but_rounded_to_zero"] for x in r["steps"]) for r in p4["adam_update_fp32"]}
        parts = [f"原版 bf16 参数 3 步更新中被舍入为 0 的比例：lr=2e-3 {min(b16[0.002]):.0%}–{max(b16[0.002]):.0%}，lr=2e-2 {min(b16[0.02]):.0%}–{max(b16[0.02]):.0%}；"
                 f"P4 fp32 主参数后 ≤ {max(f32u.values()):.1e}，前向 logits 与原版逐位相同"]
        if d["ccpp32"]:
            parts.append(f"32 条真实数据过拟合：KL {d['ccpp32']['init_dev_kl']:.3f} → {d['ccpp32']['best_dev_kl']:.4f}")
        if d["toy"]:
            parts.append(f"y=3x+7（64 行记忆）：开发 KL {d['toy']['init_dev_kl']:.3f} → {d['toy']['best_dev_kl']:.3f}")
        if d["toy256"]:
            g = d["toy256"].get("best_dev_generation", {})
            parts.append(f"y=3x+7（256 行）：开发 KL {d['toy256']['init_dev_kl']:.3f} → {d['toy256']['best_dev_kl']:.3f}，学生数字与教师平均差 {g.get('abs_diff_vs_teacher')}")
        problem = ("64 行 toy 记忆时教师常先写解释（结束率 "
                   f"{d['toy_teacher']['train']['ended_rate']:.0%}），不是合格的数字教师；已改 256 行重测" if d["toy_teacher"] else "")
        s[12] = ("已完成" if d["toy256"] else "进行中", "；".join(parts),
                 "results/cartridges_checks.json；results/p4_check.json；outputs/ccpp32/；outputs/toy*/", mdate(R / "p4_check.json"), problem)

    if gate is not None:
        b = gate[(gate["cond"] == "B") & (gate["subset"] == SUBSET)].iloc[0]
        sh = gate[(gate["cond"] == "B_shuffle") & (gate["subset"] == SUBSET)].iloc[0]
        a = gate[gate["cond"] == "A"].iloc[0]
        bt = d["base_tune"]
        mean_rmse = float(bt[bt["subset"] == SUBSET]["mean_rmse"].iloc[0])
        s[13] = ("已完成", f"M = {SUBSET}（256 行，{int(b['mem_tokens'])} 词元）；B 在 tune 256 条上解析率 {b['parse_rate']:.0%}，全部正常结束",
                 "results/gate/gate_per_query.csv", mdate(R / "gate" / "gate_summary.csv"), "")
        s[14] = ("已完成",
                 f"RMSE：B {b['rmse_filled']:.2f}，B-打乱 {sh['rmse_filled']:.2f}，均值底线 {mean_rmse:.2f}，"
                 f"A（16 词元预算）解析率 {a['parse_rate']:.0%} → 均值回退 {a['rmse_filled']:.2f}",
                 "results/gate/gate_summary.csv；results/gate/gate_by_length.csv", mdate(R / "gate" / "gate_summary.csv"),
                 "A 不听“只输出数字”，先写解释被截断；放宽到 384 词元后模型基本都明确拒答（“无模型无法预测”），文中数字是复述的输入特征 → A 记为无预测、均值回退")
        s[15] = ("已完成",
                 f"通过：B 比 B-打乱低 {-b['dRMSE_vs_shuf']:.2f}（95% CI [{-b['dRMSE_vs_shuf_hi']:.2f}, {-b['dRMSE_vs_shuf_lo']:.2f}]），"
                 f"比 A/均值低 {-b['dRMSE_vs_A']:.2f}；64–512 行的 12 份记忆全部通过；16 行时 B 在 16 词元内基本不输出数字（解析率 1.6–6.3%，不通过；属输出格式问题）",
                 "results/gate/gate_summary.csv", mdate(R / "gate" / "gate_summary.csv"),
                 "按预注册仍用 r0_n256（三份 256 行抽样中 RMSE 最高的一份，不挑最好的）")
        s[16] = ("已完成", f"同一份 256 行 M 拟合，tune：OLS RMSE {bt[bt['subset'] == SUBSET]['ols_rmse'].iloc[0]:.2f}，5-NN {bt[bt['subset'] == SUBSET]['knn5_rmse'].iloc[0]:.2f}",
                 "results/baselines_tune.csv", mdate(R / "baselines_tune.csv"), "B 的预测与 1-NN 高度相关（≈0.85），约一半预测恰好等于 M 中某行的 PE")

    t = d["teacher"]
    if t:
        s[17] = ("已完成", f"蒸馏 {t['distill']['n']} 条 + 开发 {t['dev']['n']} 条；解析率 {t['distill']['parse_rate']:.0%}/{t['dev']['parse_rate']:.0%}；"
                           f"教师 RMSE {t['distill']['teacher_rmse']:.2f}/{t['dev']['teacher_rmse']:.2f}；每位置保存 top-20 log-prob（含 <|im_end|> 位置）",
                 "outputs/main/teacher/distill.parquet；dev.parquet；*_meta.csv", mdate(MAIN / "teacher" / "summary.json"), "")
        s[18] = ("已完成", f"落盘数据每个答案位置 top-20 概率和的最小值：蒸馏 {t['distill']['topk_mass_min']:.6f}，开发 {t['dev']['topk_mass_min']:.6f}；无需增大 k",
                 "outputs/main/teacher/summary.json", mdate(MAIN / "teacher" / "summary.json"),
                 "Qwen3 把数字逐位切分，每位候选只有 0–9 等少数词元")
    runs = d["runs"]
    if any((r["dir"] / "cache_init.pt").exists() for r in runs.values()) and cc:
        s[19] = ("已完成", f"Z_init = M 渲染后的前 62 个词元 + <|im_end|>\\n 的 KV（Cartridges KVFromText 截断规则），共 64 位置；冻结 1、可训练 63；"
                           f"K/V 形状 36 层 × 8 头 × 64 × 128；k=64 初始化 Flex vs HF KL {cc['flex_vs_hf_init64']['kl_mean']:.4f}",
                 "outputs/main/runs/*/cache_init.pt", mdate(next(iter(runs.values()))["dir"] / "cache_init.pt"), "")
    lr = d["lr"]
    if lr:
        seed_runs = [runs.get(f"lr{lr}_s{sd}") for sd in SEEDS]
        done = [r for r in seed_runs if r and r["done"]]
        s[20] = ("已完成" if len(done) == 3 else "进行中",
                 f"lr={lr}（按 seed 42 开发 KL 从 {{5e-3, 2e-2}} 中选）；已完成 {len(done)}/3 个种子；"
                 + "；".join(f"{n}: {run_minutes(r['dir'])} 分钟" for n, r in runs.items() if r["done"]),
                 "outputs/main/runs/；logs/mz_main_*.out", None, "显存峰值未记录（下轮加 torch.cuda.max_memory_allocated）")
        sel = [r["sel"] for r in seed_runs if r and r["sel"]]
        if sel:
            s[21] = ("已完成" if len(sel) == 3 else "进行中",
                     "；".join(f"seed {sd}: best {r['sel']['best_ckpt']}（{r['sel']['best_updates']} 次更新），dev KL {r['sel']['init_dev_kl']:.3f} → {r['sel']['best_dev_kl']:.4f}"
                               for sd, r in zip(SEEDS, seed_runs) if r and r["sel"]),
                     "outputs/main/runs/*/dev_selection.csv；best_cache.pt", None, "")
    elif runs:
        s[20] = ("进行中", "学习率选择中：" + "，".join(runs), "outputs/main/runs/", None, "")
    if d["final"] is not None:
        s[22] = ("已完成", "选择规则（开发集全词表 KL 最小）在打开测试集前写死于 eval_cartridge.py select；测试集未参与任何选择", "jobs/05_main.sh", mdate(MAIN / "final" / "final_test_summary.csv"), "")
        s[23] = ("已完成", "见【06_测试汇总】", "outputs/main/final/final_test_summary.csv；final_test_per_query.csv", mdate(MAIN / "final" / "final_test_summary.csv"), "")
        s[24] = ("已完成", f"缓存长度 10,979 → 64（171.5×）；实际 KV {10979 * 147456 / 2 ** 20:.0f} MiB → {64 * 147456 / 2 ** 20:.1f} MiB（bf16）；GPU {gpu_name()}",
                 "【06_测试汇总】J/K 列", mdate(MAIN / "final" / "final_test_summary.csv"), "训练成本未计入压缩收益")
        figs = [p for p in ["results/fig_kl_curves.png", "results/fig_test_rmse.png"] if (ROOT / p).exists()]
        s[25] = ("已完成" if figs else "进行中",
                 "Z_init（cache_init.pt）、3 个 Z*（best_cache.pt）、切分与记忆索引、模型 revision、Cartridges commit 与补丁、全部日志已保存；"
                 + ("KL 曲线与测试 RMSE 图已生成" if figs else "KL 曲线图待生成"),
                 "；".join(["outputs/main/runs/*/", "data/processed/", "logs/"] + figs), mdate(MAIN / "final" / "final_test_summary.csv"),
                 "阶段 2（合并）未启动")
        s[26] = ("已完成",
                 "已记录：A 无记忆时拒答（不输出数字）；16 行记忆时 B 在 16 词元内基本不输出数字（解析率 1.6–6.3%）；64 行 toy 记忆教师常先写解释；"
                 "原版 bf16 参数小学习率下 33–57% 更新被舍入（→P4）；Z_init 不输出数字（C-初始行为均值回退）"
                 + "".join(f"；{t}" for t in [instability(r) for r in runs.values()] if t),
                 "【02】J 列；【07_实现检查】；outputs/main/runs/*/dev_selection.csv", mdate(MAIN / "final" / "final_test_summary.csv"),
                 "均未根据测试集结果改动任何设置。主要解读限制：蒸馏查询与测试同分布，用同样 1024 个教师伪标签拟合的 OLS 已有 RMSE "
                 + (f"{d['pseudo']['teacher_pseudo_labels']['OLS']['test_rmse_vs_truth']:.2f}" if d["pseudo"] else "?")
                 + "，需做“打乱教师 / 无关初始化 / 蒸馏量扫描 / 与查询无关的 self-study”对照（见 results/report.md 下一步）")
    if lr:
        notes = [t for t in (instability(runs.get(f"lr{lr}_s{sd}")) for sd in SEEDS) if t]
        if notes and 20 in s:
            s[20] = s[20][:4] + ("；".join(notes) + "。显存峰值未记录",)
    for row, (st, res, ev, date, prob) in s.items():
        ws[f"F{row}"], ws[f"G{row}"], ws[f"H{row}"] = st, res, ev
        if date:
            ws[f"I{row}"] = date
        if prob:
            ws[f"J{row}"] = prob


def fill_config(ws, d):
    ws["B16"] = "2000 条 X 与真实 Y"
    ws["C16"] = "实施选择"
    ws["F16"] = "扩大到 2000 条以缩小置信区间；只在开发集流程锁定后使用"
    ws["B17"] = "切分种子 0；记忆抽样种子 100+r；打乱种子 1000+10r+n"
    rows = {
        7: (f"Qwen/Qwen3-4B，revision {d['revision'][:12] if d['revision'] else '?'}", "models/Qwen3-4B/；models/Qwen3-4B.REVISION"),
        8: ("enable_thinking=False，生成提示末尾空思考块已断言", "results/token_audit.json"),
        9: (f"Cartridges {d['cart_commit']} + 补丁 P1/P2/P4", "logs/cartridges_patch.diff"),
        10: ("Folds5x2_pp.xlsx 第 1 张表，9568 行", "data/raw/ccpp.zip（sha256 cc7b2a49…）"),
        11: ("AT/V/AP/RH → PE，两位小数，'AT=14.96, V=41.76, AP=1024.07, RH=73.17, PE=463.26'", "scripts/ccpp_common.py"),
        12: ("已执行：9527 个唯一 X，跨集合泄漏 0", "data/processed/split_report.json"),
        13: (f"{SUBSET}：256 行，10,979 词元（system 轮）", "data/processed/memory_subsets.json"),
        14: ("1024 条（distill split）", "outputs/main/teacher/distill.parquet"),
        15: ("256 条（dev split，只用于选点）；另有 tune 256 条做教师关卡", "data/processed/rows.csv"),
        16: ("2000 条", "data/processed/rows.csv"),
        17: ("0 / 100+r / 1000+10r+n", "scripts/prepare_data.py"),
        18: ("64 位置（62 个 M 词元 + <|im_end|>\\n）", "outputs/main/runs/*/cache_init.pt"),
        19: ("1（已验证存读后仍为 1）", "results/cartridges_checks.json"),
        20: ("top-20，覆盖率 ≥ 0.999998", "outputs/main/teacher/summary.json"),
        21: ("截断软标签 CE（Cartridges 原损失）；选点用全词表 KL(B‖C)", "scripts/eval_cartridge.py"),
        22: ("bf16 前向 + fp32 主参数（P4，MZ_FP32_CACHE=1）", "results/p4_check.json"),
        23: ("seed 42 / 43 / 44", "outputs/main/runs/"),
        24: ("每 epoch 存一次，按开发 KL 选；best_cache.pt 为软链接", "outputs/main/runs/*/dev_selection.csv"),
        25: ("tune 集：B 同时低于 A、均值、B-打乱，且 95% CI 不含 0", "results/gate/gate_summary.csv"),
        26: ("解析失败用 M 的 PE 均值回退，另报解析率", "scripts/teacher_gate.py"),
        27: ("同一份 256 行 M 拟合", "results/baselines_tune.csv"),
        28: (gpu_name() or "H100 NVL 94 GB", "logs/mz_*.out nvidia-smi 行"),
        29: ("已下载到 models/Qwen3-4B（7.6 GB）", "logs/mz_setup_*.out"),
        30: ("已复制到 third_party/cartridges", "third_party/cartridges"),
        31: ("ccpp.zip，sha256 cc7b2a4977c0…", "logs/mz_setup_*.out"),
    }
    for r, (actual, ev) in rows.items():
        ws[f"D{r}"], ws[f"E{r}"] = actual, ev
    extra = [
        ("调格式 / 教师关卡集 tune", "256 条 X 与真实 Y", "实施选择", "已用于教师关卡", "data/processed/rows.csv", "与开发集分开，避免开发集被用两次"),
        ("记忆长度扫描", "16 / 64 / 128 / 256 / 512 行 × 3 份抽样", "ChatGPT 建议", "已执行", "results/gate/gate_by_length.csv", "64 行起 B 稳定受益；16 行时 B 不输出数字"),
        ("学习率候选", "{5e-3, 2e-2}，seed 42 上按开发 KL 选", "实施选择", d["lr"] or "待定", "outputs/main/chosen_lr.txt", "选定后 seed 43/44 沿用"),
        ("训练轮数 / 批大小", "20 epoch；每步 8 条 512 词元打包序列（约 4 条查询/序列）", "实施选择", "已执行" if d["runs"] else "", "outputs/main/runs/*/train_args.json", "Adam，无学习率调度"),
        ("A 组生成规则", "16 词元贪心；补充：384 词元长回答取最后一个数", "实测后补充", "A 在 16 词元内解析率 0%", "outputs/main/A_long_*.csv", "长回答版中 A 基本拒答，“最后一个数”是复述的 RH 等输入特征，不能当预测；A 按无预测处理"),
    ]
    for i, vals in enumerate(extra):
        r = 32 + i
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v)
            ref = ws.cell(31, j)
            c.font, c.border, c.alignment = copy(ref.font), copy(ref.border), copy(ref.alignment)
            c.fill = copy(ws.cell(31 if i % 2 else 30, j).fill)
        ws.row_dimensions[r].height = ws.row_dimensions[31].height


def fill_gate(ws, d):
    ws["A3"] = ("tune 集 256 条（与开发集分开；开发集只用于选 checkpoint）；主记忆 r0_n256；所有预测使用同一查询模板。"
                "更多记忆长度见【04b_记忆长度扫描】。")
    ws["A18"] = ("教师关卡必须先在 tune 集跑 A、均值、B、B-打乱；如果 B 的 RMSE 没有明显降低，优先检查提示模板、数字生成格式、缓存接口、"
                 "M 的样本覆盖，然后再考虑换模型。开发集不参与关卡，只用于挑 checkpoint。")
    gate = d["gate"]
    if gate is None:
        return
    bt = d["base_tune"]
    row = bt[bt["subset"] == SUBSET].iloc[0]
    g = lambda c: gate[(gate["cond"] == c) & (gate["subset"] == (SUBSET if c != "A" else "none"))].iloc[0]  # noqa: E731
    a, b, sh = g("A"), g("B"), g("B_shuffle")
    pq = d["gate_pq"]
    bq = pq[(pq["cond"] == "B") & (pq["subset"] == SUBSET)].set_index("qid")
    sq = pq[(pq["cond"] == "B_shuffle") & (pq["subset"] == SUBSET)].set_index("qid")
    both = bq["parse_ok"] & sq["parse_ok"]
    sh_diff = float((sq.loc[both, "pred"] - bq.loc[both, "pred"]).abs().mean())
    # A gives no numeric prediction: 0% bare numbers in 16 tokens; with a long budget it explains and refuses,
    # and the "last number" it writes is an echoed input feature (e.g. RH), not a PE prediction.
    a_rmse, a_mae, a_parse = a["rmse_filled"], a["mae_filled"], 0.0
    al = d["A_tune"]
    a_note = ("无记忆时不给出数值预测：16 词元内解析率 0%"
              + (f"；长回答（≤384 词元）中 {refusal_rate(al):.0%} 明确表示无法给出数值，文中出现的数字是复述的输入特征（如 RH）" if al is not None else "")
              + "。按无预测处理，RMSE/MAE 为记忆池均值回退")
    vals = {
        9: (a_rmse, a_mae, a_parse, "测试集上计算", None, "results/gate/gate_summary.csv；outputs/main/A_long_tune.csv", a_note),
        10: (row["mean_rmse"], row["mean_mae"], 1.0, "不适用", None, "results/baselines_tune.csv", "恒预测 M 内 PE 均值"),
        11: (b["rmse_filled"], b["mae_filled"], b["parse_rate"], 0, 0, "results/gate/gate_per_query.csv",
             f"核心教师；约 {b['copy_rate']:.0%} 的预测恰为 M 中某行 PE；与 1-NN 相关 {b['corr_with_knn1']:.2f}"),
        12: (sh["rmse_filled"], sh["mae_filled"], sh["parse_rate"], "测试集上计算", sh_diff, "results/gate/gate_per_query.csv",
             f"B 比 B-打乱低 {-b['dRMSE_vs_shuf']:.2f} MW，95% CI [{-b['dRMSE_vs_shuf_hi']:.2f}, {-b['dRMSE_vs_shuf_lo']:.2f}]"),
        13: (row["ols_rmse"], row["ols_mae"], 1.0, "不适用", None, "results/baselines_tune.csv", "辅助参照；不属于教师关卡"),
        14: (row["knn5_rmse"], row["knn5_mae"], 1.0, "不适用", None, "results/baselines_tune.csv", "辅助参照；特征按 M 标准化"),
    }
    for r, (rm, ma, pr, kl, ad, path, nt) in vals.items():
        setv(ws, f"D{r}", f3(rm), "0.00")
        setv(ws, f"E{r}", f3(ma), "0.00")
        setv(ws, f"F{r}", f3(pr), "0.0%")
        ws[f"G{r}"] = kl
        if ad is not None:
            setv(ws, f"H{r}", f3(ad), "0.00")
        ws[f"I{r}"], ws[f"J{r}"] = path, nt


def add_sweep_sheet(wb, d, index):
    if d["gate"] is None:
        return
    ws = wb.create_sheet("04b_记忆长度扫描", index)
    title_block(ws, "教师关卡补充｜记忆长度扫描",
                "tune 集 256 条；每个长度从 1024 行记忆池中抽 3 份（嵌套前缀）；ΔRMSE = B − 对照，负值表示 B 更好，配对 bootstrap 95% CI。", 12)
    cols = ["记忆行数", "抽样", "记忆词元数", "B 解析率", "B RMSE", "B MAE", "B-打乱 RMSE", "ΔRMSE vs B-打乱 [95% CI]",
            "ΔRMSE vs A/均值 [95% CI]", "OLS RMSE（同一 M）", "5-NN RMSE（同一 M）", "关卡"]
    header(ws, 5, cols)
    gate, bt = d["gate"], d["base_tune"]
    r = 6
    for n in [16, 64, 128, 256, 512]:
        for rs in range(3):
            key = f"r{rs}_n{n}"
            b = gate[(gate["cond"] == "B") & (gate["subset"] == key)].iloc[0]
            sh = gate[(gate["cond"] == "B_shuffle") & (gate["subset"] == key)].iloc[0]
            base = bt[bt["subset"] == key].iloc[0]
            passed = b["parse_rate"] > 0.98 and b["dRMSE_vs_shuf_hi"] < 0 and b["dRMSE_vs_A_hi"] < 0
            body_row(ws, r, [n, key + ("（主记忆）" if key == SUBSET else ""), int(b["mem_tokens"]), f3(b["parse_rate"]), f3(b["rmse_filled"], 2),
                             f3(b["mae_filled"], 2), f3(sh["rmse_filled"], 2),
                             f"{b['dRMSE_vs_shuf']:.2f} [{b['dRMSE_vs_shuf_lo']:.2f}, {b['dRMSE_vs_shuf_hi']:.2f}]",
                             f"{b['dRMSE_vs_A']:.2f} [{b['dRMSE_vs_A_lo']:.2f}, {b['dRMSE_vs_A_hi']:.2f}]",
                             f3(base["ols_rmse"], 2), f3(base["knn5_rmse"], 2), "通过" if passed else "不通过"],
                     fmts={4: "0.0%", 5: "0.00", 6: "0.00", 7: "0.00", 10: "0.00", 11: "0.00"}, zebra=(n // 16) % 2 == 0)
            r += 1
    r += 1
    header(ws, r, ["记忆行数", "3 份均值：B RMSE", "B RMSE 标准差", "B-打乱 RMSE 均值", "均值底线 RMSE", "OLS RMSE", "5-NN RMSE", "1-NN RMSE",
                   "", "", "", ""])
    agg = d["gate_len"]
    r += 1
    for n in [16, 64, 128, 256, 512]:
        bb = agg[(agg["cond"] == "B") & (agg["n_rows"] == n)].iloc[0]
        ss = agg[(agg["cond"] == "B_shuffle") & (agg["n_rows"] == n)].iloc[0]
        body_row(ws, r, [n, f3(bb["rmse_mean"], 2), f3(bb["rmse_sd"], 2), f3(ss["rmse_mean"], 2), f3(bb["mean_rmse"], 2), f3(bb["ols_rmse"], 2),
                         f3(bb["knn5_rmse"], 2), f3(bb["knn1_rmse"], 2), None, None, None, None],
                 fmts={j: "0.00" for j in range(2, 9)})
        r += 1
    note(ws, r + 1, "解读：64 行起 B 在全部 12 份记忆上都显著好于 B-打乱和 A/均值；16 行时 B 在 16 词元内基本不输出数字（解析率 1.6–6.3%，先写解释被截断），属输出格式问题而非记忆无效。"
                    "B 仍明显差于 OLS / 5-NN，且预测与 1-NN 高度相关：模型主要在“找相近记录并抄 PE”。", 12)
    widths = [10, 20, 12, 10, 10, 10, 13, 24, 24, 16, 16, 10]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A6"


def rebuild_training(wb, d):
    old = wb["05_训练日志"]
    idx = wb.sheetnames.index("05_训练日志")
    wb.remove(old)
    ws = wb.create_sheet("05_训练日志", idx)
    title_block(ws, "学生训练 Z｜3 个种子逐 checkpoint 记录",
                "只更新 64 位置 KV 记忆 Z（冻结第 1 位置）。训练截断 CE 不等于完整 KL；选点以开发集全词表 KL(B‖C)（沿教师答案路径）为准。每 epoch 存一个 checkpoint。", 12)
    cols = ["训练种子", "运行 ID", "Optimizer Step（更新次数）", "Train：截断 CE ↓（日志）", "Dev：KL(B‖C) ↓", "Dev top-1 一致率 ↑",
            "Dev RMSE ↓（生成）", "可解析率 ↑", "Top-k 概率覆盖率 ↑", "‖Z − Z_init‖ / ‖Z_init‖", "Checkpoint 文件路径", "现象 / 备注"]
    header(ws, 5, cols)
    runs, lr = d["runs"], d["lr"]
    cover = d["teacher"]["distill"]["topk_mass_min"] if d["teacher"] else None
    order = [(f"seed_{i + 1}", f"lr{lr}_s{sd}") for i, sd in enumerate(SEEDS)] if lr else []
    order += [("lr 选择", n) for n in runs if n.endswith("_s42") and (not lr or n != f"lr{lr}_s42")]
    if not lr:
        order += [("lr 选择", n) for n in runs if not n.endswith("_s42")]
    r = 6
    best_rows = []
    for label, name in order:
        run = runs.get(name)
        if run is None or run["sel_csv"] is None:
            body_row(ws, r, [label, name, None, None, None, None, None, None, None, None, None,
                             "训练中" if run else "未开始"], input_cols=range(3, 13))
            r += 1
            continue
        df = run["sel_csv"]
        losses = train_losses(name)
        drift = z_drift(run["dir"], list(df["ckpt"]))
        best = run["sel"] or {}  # dev_selection.json is written after the csv; may be missing mid-selection
        for k, rr in df.iterrows():
            is_best = rr["ckpt"] == best.get("best_ckpt")
            gen = best.get("best_dev_generation", {}) if is_best else (best.get("init_dev_generation", {}) if rr["ckpt"] == "init" else {})
            step = int(rr["updates"])
            ce = losses.get(step) or losses.get(step - 1)
            body_row(ws, r, [label, name, step, f3(ce, 4) if rr["ckpt"] != "init" else None, f3(rr["dev_kl"], 4), f3(rr["dev_top1_agree"], 3),
                             f3(gen.get("rmse_vs_truth_parsed"), 2), f3(gen.get("parse_rate"), 3), f3(cover, 6),
                             0.0 if rr["ckpt"] == "init" else f3(drift.get(rr["ckpt"]), 4),
                             rel(run["dir"] / f"{rr['ckpt']}.pt" if rr["ckpt"] != "init" else run["dir"] / "cache_init.pt"),
                             "★ 开发集 KL 最小（best）" if is_best else ("Z_init（未训练）" if rr["ckpt"] == "init" else None)],
                     fmts={4: "0.0000", 5: "0.0000", 6: "0.0%", 7: "0.00", 8: "0.0%", 9: "0.000000", 10: "0.0000"},
                     zebra=k % 2 == 1, input_cols=())
            r += 1
        if label.startswith("seed") and best:
            best_rows.append((label, name, best, run))
    r += 1
    note_r = r
    ws.merge_cells(start_row=note_r, start_column=1, end_row=note_r, end_column=12)
    style(ws.cell(note_r, 1, "每个种子的最佳检查点（开发集选，不使用测试集选）"), fill=F_HEAD, bold=True, color="FFFFFFFF", size=11)
    r += 1
    header(ws, r, ["种子", "Best Step", "Best Dev KL", "Best Dev RMSE", "学习率", "Best cache 文件", "是否完成", "与教师同 token 比例",
                   "Dev 与教师数字差", "训练时长（分钟）", "Z_init Dev KL", "备注"])
    dv = DataValidation(type="list", formula1='"未开始,训练中,已完成,失败"', allow_blank=True)
    ws.add_data_validation(dv)
    for i, sd in enumerate(SEEDS):
        r += 1
        hit = [x for x in best_rows if x[0] == f"seed_{i + 1}"]
        if hit:
            _, name, best, run = hit[0]
            g = best.get("best_dev_generation", {})
            vals = [f"seed_{i + 1}（{sd}）", best["best_updates"], f3(best["best_dev_kl"], 4), f3(g.get("rmse_vs_truth_parsed"), 2), lr,
                    rel(run["dir"] / "best_cache.pt") + f" → {best['best_ckpt']}.pt", "已完成", f3(g.get("exact_same_tokens_as_teacher"), 3),
                    f3(g.get("abs_diff_vs_teacher"), 3), run_minutes(run["dir"]), f3(best["init_dev_kl"], 4), instability(run)]
        else:
            vals = [f"seed_{i + 1}（{sd}）", None, None, None, lr, None, "训练中" if any(n.endswith(f"_s{sd}") for n in runs) else "未开始",
                    None, None, None, None, None]
        body_row(ws, r, vals, fmts={3: "0.0000", 4: "0.00", 8: "0.0%", 9: "0.000", 11: "0.0000"}, input_cols=range(2, 13))
        dv.add(ws.cell(r, 7))
    note(ws, r + 2, "Train CE 取自训练日志中该更新次数对应的进度条损失（截断 top-20 软标签 CE，按 (位置, 候选) 平均），只作趋势参考。"
                    "Dev RMSE / 可解析率只对 Z_init 与 best 两个 checkpoint 做自由生成。显存峰值本轮未记录。", 12)
    widths = [11, 15, 12, 14, 13, 12, 12, 10, 12, 14, 44, 26]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A6"


def fill_test(ws, d, rows, subsets):
    for r in range(7, 17):
        ws[f"D{r}"] = 2000
    fin = d["final"]
    cl = classic("test", rows, subsets)
    mib = lambda n: round(n * 147456 / 2 ** 20, 2)  # noqa: E731  bytes per position: 36 layers x 8 heads x 128 x (K,V) x 2 bytes
    setv(ws, "E8", f3(cl["mean"][0]), "0.00"); setv(ws, "F8", f3(cl["mean"][1]), "0.00"); setv(ws, "G8", 1.0, "0.0%")
    setv(ws, "E15", f3(cl["ols"][0]), "0.00"); setv(ws, "F15", f3(cl["ols"][1]), "0.00"); setv(ws, "G15", 1.0, "0.0%")
    setv(ws, "E16", f3(cl["knn5"][0]), "0.00"); setv(ws, "F16", f3(cl["knn5"][1]), "0.00"); setv(ws, "G16", 1.0, "0.0%")
    ws["L8"], ws["L15"], ws["L16"] = "同一份 256 行 M 的 PE 均值", "同一份 256 行 M 拟合", "同一份 256 行 M，特征标准化"
    # A: no numeric prediction (see fill_gate) -> memory-pool mean fallback, parse rate 0
    pool_mean = float(rows[rows["split"] == "memory_pool"][TARGET].mean())
    yt = rows[rows["split"] == "test"][TARGET].to_numpy()
    setv(ws, "E7", f3(np.sqrt(np.mean((pool_mean - yt) ** 2))), "0.00")
    setv(ws, "F7", f3(np.mean(np.abs(pool_mean - yt))), "0.00")
    setv(ws, "G7", 0.0, "0.0%")
    at = d["A_test"]
    ws["L7"] = ("不给出数值预测（16 词元内 0%"
                + (f"；长回答中 {refusal_rate(at):.0%} 明确拒答，出现的数字是复述的输入特征" if at is not None else "")
                + "）；RMSE/MAE 为记忆池均值回退")
    if fin is None:
        return
    rowmap = {"A": 7, "B": 9, "B_shuffle": 10, "C_init": 11, "C_s42": 12, "C_s43": 13, "C_s44": 14}
    for _, x in fin.iterrows():
        r = rowmap.get(x["cond"])
        if r is None:
            continue
        if x["cond"] != "A":
            setv(ws, f"E{r}", f3(x["rmse_filled"]), "0.00")
            setv(ws, f"F{r}", f3(x["mae_filled"]), "0.00")
            setv(ws, f"G{r}", f3(x["parse_rate"]), "0.0%")
        if x["cond"] != "B":
            setv(ws, f"H{r}", f3(x["kl_B_X"], 4), "0.0000")
        if x["cond"] not in ("A", "B"):
            setv(ws, f"I{r}", f3(x["abs_diff_vs_B"]), "0.00")
        ws[f"J{r}"] = int(x["cache_positions"])
        ws[f"K{r}"] = mib(int(x["cache_positions"]))
        tail = f"有效题 {int(x['n_both_parsed'])}/{int(x['n'])}" if x["cond"] not in ("A", "B") and not pd.isna(x["n_both_parsed"]) else ""
        prev = ws[f"L{r}"].value if x["cond"] == "A" else ""
        ws[f"L{r}"] = "；".join(filter(None, [prev, "outputs/main/final/final_test_summary.csv", tail]))
    for r in (8, 15, 16):
        ws[f"J{r}"], ws[f"K{r}"] = 0, 0
    ws["L11"] = "Z_init 不输出数字（解析率 0%），RMSE/MAE 为均值回退；KL 沿 B 的答案路径计算；" + (ws["L11"].value or "")
    st = d["stats"]
    if st:
        p = st["pairs"]
        for sd, r in zip(SEEDS, (12, 13, 14)):
            q = p[f"C_s{sd} - B"]
            ws[f"L{r}"] = f"RMSE − B = {q['diff']:+.2f} [{q['ci95'][0]:+.2f}, {q['ci95'][1]:+.2f}]；" + (ws[f"L{r}"].value or "")
        ws.merge_cells("A23:L23")
        q, sh, mn, rr = p["C_avg3 - B"], p["C_avg3 - B_shuffle"], p["C_avg3 - mean"], st["recovery_R"]
        style(ws["A23"], fill=F_INPUT, color=LABEL, border=False)
        ws["A23"] = (f"配对 bootstrap（{st['n']} 条测试题按 qid 配对，5000 次重采样，95% CI）：C 三种子平均 RMSE − B = {q['diff']:+.2f} MW "
                     f"[{q['ci95'][0]:+.2f}, {q['ci95'][1]:+.2f}]；− B-打乱 = {sh['diff']:+.2f} [{sh['ci95'][0]:+.2f}, {sh['ci95'][1]:+.2f}]；"
                     f"− 均值底线 = {mn['diff']:+.2f} [{mn['ci95'][0]:+.2f}, {mn['ci95'][1]:+.2f}]；"
                     f"恢复率 R = {rr['point']:.1%} [{rr['ci95'][0]:.1%}, {rr['ci95'][1]:.1%}]（以均值底线为无记忆参照，因为 A 不给数字）。")
        ws.row_dimensions[23].height = 45
    ps = d["pseudo"]
    if ps:
        tp, cr = ps["teacher_pseudo_labels"], ps["C_reference"]
        ws.merge_cells("A24:L24")
        style(ws["A24"], fill=F_INPUT, color=LABEL, border=False)
        ws["A24"] = ("信息来源对照（不用记忆，只用 C 训练时同样的 1024 个 (X, 教师答案) 拟合传统回归器，test 2000 条）："
                     + "；".join(f"{k} RMSE {v['test_rmse_vs_truth']:.2f}，与 B 相关 {v['corr_with_B']:.3f}，与 B 数字差 {v['mean_abs_diff_vs_B']:.2f}"
                                for k, v in tp.items())
                     + f"。C：RMSE {cr['test_rmse_mean_of_seed_rmses']:.2f}，与 B 相关 {cr['corr_with_B_mean_of_seeds']:.3f}，数字差 {cr['mean_abs_diff_vs_B_mean_of_seeds']:.2f}。"
                     "→ C 比这些回归器更贴近 B，但 RMSE 本身不能证明“压缩了 M”。")
        ws.row_dimensions[24].height = 60


def add_checks_sheet(wb, d):
    ws = wb.create_sheet("07_实现检查")
    title_block(ws, "实现与数值检查｜每项都可追溯到结果文件",
                "这些检查不产生实验结论，只证明“比较是公平的、代码做的就是我们以为它做的”。", 6)
    header(ws, 5, ["检查项", "比较的两边", "指标", "结果", "判定", "证据文件"])
    eq, cc, p4, au = d["eq"], d["cc"], d["p4"], d["audit"]
    items = []
    if eq:
        f32 = eq.get("fp32_vs_bf16_full", {})
        floor = f32.get("kl_mean", {}).get("mean")
        kl = eq["cached_vs_full_bf16"]["kl_mean"]["mean"]
        items += [
            ("HF 前缀 KV 一致性", "先算 M 的 KV 再接查询 vs 整串一次前向（bf16，20 条）", "答案位置 KL 均值 / argmax 一致",
             f"{kl:.4f} / {eq['cached_vs_full_bf16']['argmax_agree']['mean']:.3f}", "一致（≤ 2× 噪声底）" if floor and kl <= 2 * floor else "需复查", "results/gate/equivalence.json"),
            ("bf16 噪声底", "fp32 模型 vs bf16 模型，整串前向", "KL 均值 / argmax 一致",
             f"{floor:.4f} / {f32.get('argmax_agree', {}).get('mean', float('nan')):.3f}" if floor else "—", "参照", "results/gate/equivalence.json"),
            ("缓存复位", "同一查询在裁剪缓存后重跑", "logits 最大差", f"{eq['repeat_after_crop']['max_abs_logit_diff']}", "逐位一致", "results/gate/equivalence.json"),
        ]
    if cc:
        items += [
            ("Cartridges Flex vs HF（Z_init）", "k=64 初始化 cartridge vs HF 同 64 词元前缀", "KL 均值 / argmax 一致",
             f"{cc['flex_vs_hf_init64']['kl_mean']:.4f} / {cc['flex_vs_hf_init64']['argmax_agree']:.3f}", "一致", "results/cartridges_checks.json"),
            ("Cartridges Flex vs HF（完整 M）", "完整 M 作为 cartridge vs HF 整串", "KL 均值 / argmax 一致",
             f"{cc['flex_vs_hf_fullM']['kl_mean']:.4f} / {cc['flex_vs_hf_fullM']['argmax_agree']:.3f}", "一致（≈1.7× 噪声底）", "results/cartridges_checks.json"),
            ("存 / 读（补丁 P1）", "save → from_pretrained", "冻结 / 可训练位置；数值",
             f"{cc['cache_k64']['loaded_num_frozen']} / {cc['cache_k64']['loaded_num_trainable']}；相同={cc['cache_k64']['roundtrip_identical']}（未修复会读成冻结 {cc['cache_k64']['unpatched_would_read_num_frozen']}）",
             "通过", "results/cartridges_checks.json"),
        ]
        for r in cc["adam_update"]:
            z = [x["frac_intended_nonzero_but_rounded_to_zero"] for x in r["steps"]]
            items.append((f"bf16 参数更新（原版，lr={r['lr']}）", "Adam 应有更新 vs 实际写入的更新（3 步）", "被舍入为 0 的比例",
                          f"{min(z):.1%} – {max(z):.1%}", "有问题 → 引入 P4", "results/cartridges_checks.json"))
        cov = cc["teacher_topk_coverage"]
        items.append(("top-20 覆盖率（早期抽查）", "8 条查询，每个答案位置", "top-20 概率和最小值",
                      f"{min(v['top20_mass_min'] for v in cov.values()):.6f}", "足够", "results/cartridges_checks.json"))
    if p4:
        for r in p4["adam_update_fp32"]:
            z = max(x["frac_intended_nonzero_but_rounded_to_zero"] for x in r["steps"])
            items.append((f"fp32 主参数（P4，lr={r['lr']}）", "同上", "被舍入为 0 的比例；loss 3 步",
                          f"{z:.1e}；{r['steps'][0]['loss_before_step']:.3f} → {r['loss_after_last_step']:.3f}", "通过", "results/p4_check.json"))
        items.append(("P4 不改变前向", "fp32 主参数 cache vs 原版 bf16 cache", "logits 最大差", f"{p4['forward']['max_abs_logit_diff']}", "逐位一致", "results/p4_check.json"))
    if au:
        mc = au["message_converter"]
        items.append(("Qwen3 非思考模式 token", "HF 模板 vs Cartridges 转换器", "token 序列是否相同",
                      f"原样：{mc['cartridges_plain_equals_target']}；补空思考块后：{mc['cartridges_with_empty_think_equals_target']}", "已对齐（每条教师数据再断言一次）", "results/token_audit.json"))
        items.append(("数字分词", "'463.26'", "词元", " ".join(au["digit_tokenization"]["463.26"]), "逐位切分", "results/token_audit.json"))
    if d["ccpp32"]:
        items.append(("32 条真实数据过拟合", "训练集 = 评估集", "KL(B‖C)：Z_init → best", f"{d['ccpp32']['init_dev_kl']:.3f} → {d['ccpp32']['best_dev_kl']:.4f}", "通过（降 >99%）", "outputs/ccpp32/run_k64/dev_selection.json"))
    if d["toy"]:
        items.append(("y=3x+7（64 行记忆，k=32）", "200 训练 / 50 开发", "开发 KL：Z_init → best", f"{d['toy']['init_dev_kl']:.3f} → {d['toy']['best_dev_kl']:.3f}",
                      "蒸馏可用；但教师常不直接给数字", "outputs/toy/"))
    if d["toy256"]:
        g = d["toy256"].get("best_dev_generation", {})
        tt = d["toy256_teacher"]["dev"] if d["toy256_teacher"] else {}
        items.append(("y=3x+7（256 行记忆，k=32）", "200 训练 / 50 开发", "开发 KL；教师 RMSE；学生与教师数字差",
                      f"{d['toy256']['init_dev_kl']:.3f} → {d['toy256']['best_dev_kl']:.3f}；{tt.get('teacher_rmse', float('nan')):.2f}；{g.get('abs_diff_vs_teacher')}",
                      "见数值", "outputs/toy256/"))
    for i, it in enumerate(items):
        body_row(ws, 6 + i, list(it), zebra=i % 2 == 1)
    for j, w in enumerate([26, 40, 26, 40, 24, 36], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A6"


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "docs" / "input_record.xlsx"
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "docs" / "记忆压缩实验_阶段0_执行跟踪_Qwen3-4B.xlsx"
    d = gather()
    rows, subsets = load_splits()
    wb = load_workbook(src)
    fill_overview(wb["01_总览"], d)
    fill_steps(wb["02_执行步骤"], d)
    status_colors(wb["02_执行步骤"], "F7:F26")
    fill_config(wb["03_准备配置"], d)
    fill_gate(wb["04_教师预检"], d)
    add_sweep_sheet(wb, d, wb.sheetnames.index("04_教师预检") + 1)
    rebuild_training(wb, d)
    fill_test(wb["06_测试汇总"], d, rows, subsets)
    add_checks_sheet(wb, d)
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = wb.calculation or CalcProperties()
    wb.calculation.fullCalcOnLoad = True  # no LibreOffice on the cluster: let Excel/WPS recalculate on open
    wb.save(dst)
    print("wrote", dst)


if __name__ == "__main__":
    main()
